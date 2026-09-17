"""Tests for the unified file-ingestion tool + terminal /file UX."""

import os

import pytest


def _make_pdf(text: str) -> bytes:
    """Build a minimal valid single-page PDF containing ``text``."""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
    ]
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    pdf = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(pdf)
    pdf += f"xref\n0 {len(objs) + 1}\n".encode()
    pdf += b"0000000000 65535 f \n"
    for off in offsets:
        pdf += f"{off:010d} 00000 n \n".encode()
    pdf += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return pdf


@pytest.fixture
def tmpfiles(tmp_path):
    """Create fixture files: text, pdf, docx, xlsx, pptx, image, audio."""
    text = tmp_path / "notes.txt"
    text.write_text("Homework notes line one\nline two", encoding="utf-8")
    pdf = tmp_path / "hw.pdf"
    pdf.write_bytes(_make_pdf("Homework PDF body text"))
    img = tmp_path / "page.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
    audio = tmp_path / "lecture.mp3"
    audio.write_bytes(b"ID3\x04\x00\x00\x00\x00\x00\x00")
    docx = tmp_path / "essay.docx"
    import docx as _docx

    d = _docx.Document()
    d.add_paragraph("Word essay content")
    d.save(str(docx))
    xlsx = tmp_path / "data.xlsx"
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a", "b"])
    ws.append([1, 2])
    wb.save(str(xlsx))
    pptx = tmp_path / "slides.pptx"
    import pptx as _pptx

    prs = _pptx.Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Slide one"
    prs.save(str(pptx))
    return {
        "text": text, "pdf": pdf, "img": img, "audio": audio,
        "docx": docx, "xlsx": xlsx, "pptx": pptx,
    }


class TestIngestDispatch:
    def test_text(self, tmpfiles):
        from tools.ingest_tools import ingest_file

        out = ingest_file(str(tmpfiles["text"]), summarize=False, max_chars=100)
        assert "Homework notes line one" in out

    def test_pdf(self, tmpfiles):
        from tools.ingest_tools import ingest_file

        out = ingest_file(str(tmpfiles["pdf"]), summarize=False, max_chars=200)
        assert "Homework PDF body text" in out

    def test_image_uses_vision(self, tmpfiles, monkeypatch):
        from tools import ingest_tools

        monkeypatch.setattr(
            "tools.vision_tools.ocr_document", lambda path: "OCR TEXT FROM IMAGE"
        )
        out = ingest_tools.ingest_file(str(tmpfiles["img"]), summarize=False)
        assert "OCR TEXT FROM IMAGE" in out

    def test_audio_transcribes(self, tmpfiles, monkeypatch):
        from tools import ingest_tools

        monkeypatch.setattr("stt.transcribe_file", lambda path: "Transcribed lecture text")
        out = ingest_tools.ingest_file(str(tmpfiles["audio"]), summarize=False)
        assert "Transcribed lecture text" in out

    def test_docx(self, tmpfiles):
        from tools.ingest_tools import ingest_file

        out = ingest_file(str(tmpfiles["docx"]), summarize=False, max_chars=200)
        assert "Word essay content" in out

    def test_xlsx(self, tmpfiles):
        from tools.ingest_tools import ingest_file

        out = ingest_file(str(tmpfiles["xlsx"]), summarize=False, max_chars=200)
        assert "a | b" in out and "1 | 2" in out

    def test_pptx(self, tmpfiles):
        from tools.ingest_tools import ingest_file

        out = ingest_file(str(tmpfiles["pptx"]), summarize=False, max_chars=200)
        assert "Slide one" in out

    def test_missing_file(self):
        from tools.ingest_tools import ingest_file

        assert "File not found" in ingest_file("/no/such/file.pdf")

    def test_unsupported_extension(self, tmp_path):
        from tools.ingest_tools import ingest_file

        f = tmp_path / "blob.bin"
        f.write_bytes(b"\x00\x01\x02\x03")
        out = ingest_file(str(f), summarize=False)
        assert "Unsupported file type" in out


class TestCapSummarize:
    def test_large_file_capped_and_summarized(self, tmp_path, monkeypatch):
        from tools import ingest_tools

        big = tmp_path / "big.txt"
        big.write_text("word " * 5000, encoding="utf-8")  # ~25k chars
        monkeypatch.setattr(ingest_tools, "_summarize", lambda text: "SUMMARY OF REMAINDER")
        out = ingest_tools.ingest_file(str(big), max_chars=1000)
        assert "SUMMARY OF REMAINDER" in out
        assert "has" in out and "chars" in out
        assert len(out) < 3000  # bounded

    def test_small_file_not_summarized(self, tmp_path, monkeypatch):
        from tools import ingest_tools

        small = tmp_path / "small.txt"
        small.write_text("short", encoding="utf-8")
        called = []
        monkeypatch.setattr(ingest_tools, "_summarize", lambda text: called.append(1) or "S")
        out = ingest_tools.ingest_file(str(small))
        assert "short" in out
        assert not called

    def test_size_guard(self, tmp_path, monkeypatch):
        from tools import ingest_tools

        big = tmp_path / "huge.bin"
        big.write_bytes(b"\x00" * 1024)
        monkeypatch.setattr(ingest_tools, "INGEST_MAX_BYTES", 100)
        out = ingest_tools.ingest_file(str(big), summarize=False)
        assert "too large" in out.lower()


class TestTerminalUX:
    def test_looks_like_file_path(self):
        from terminal import _looks_like_file_path

        p = os.path.abspath(__file__)
        assert _looks_like_file_path(p) == p
        assert _looks_like_file_path(f"~/Jarvis/{os.path.relpath(p, os.path.expanduser('~/Jarvis'))}")
        assert _looks_like_file_path(f"file://{p}") == p
        assert _looks_like_file_path("hello world this is a message") is None
        assert _looks_like_file_path("/no/such/path/x.pdf") is None

    def test_file_command_routes_true(self, tmpfiles, monkeypatch):
        from terminal import handle_local_command

        calls = []
        monkeypatch.setattr(
            "brain.ingest_file_into_context",
            lambda session, path: calls.append(path) or "ingested content",
        )
        assert handle_local_command(f"/file {tmpfiles['text']}") is True
        assert calls == [str(tmpfiles["text"])]

    def test_bare_path_auto_detects(self, tmpfiles, monkeypatch):
        from terminal import handle_local_command

        calls = []
        monkeypatch.setattr(
            "brain.ingest_file_into_context",
            lambda session, path: calls.append(path) or "ingested content",
        )
        assert handle_local_command(str(tmpfiles["text"])) is True
        assert calls == [str(tmpfiles["text"])]

    def test_normal_message_not_consumed(self):
        from terminal import handle_local_command

        assert handle_local_command("what is the weather?") is False


class TestContextHelper:
    def test_ingest_file_into_context_appends(self, tmpfiles, monkeypatch):
        import brain

        result = brain.ingest_file_into_context("sess_test", str(tmpfiles["text"]))
        assert "Homework notes line one" in result
        # Appended to the in-memory conversation buffer.
        assert any("ingested file" in m["content"] for m in brain.conversation)
